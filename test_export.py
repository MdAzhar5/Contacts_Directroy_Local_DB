from __future__ import annotations

import csv
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import db
from desktop_app import ContactDirectoryDesktop

filters = {"field": "" for field in ["q", "industry", "city", "state", "country", "category", "source", "source_type", "source_category", "source_file"]}
filters["industry"] = "Medical NPI"
_, expected = db.query_contacts(filters, page=1, per_page=1)
assert expected > 0

with tempfile.TemporaryDirectory() as tmp:
    tmp_path = Path(tmp)
    app = ContactDirectoryDesktop()
    try:
        csv_path = tmp_path / "filtered.csv"
        xlsx_path = tmp_path / "filtered.xlsx"
        _, csv_count, csv_type = app._export_rows(csv_path, "csv", filters)
        _, xlsx_count, xlsx_type = app._export_rows(xlsx_path, "xlsx", filters)
    finally:
        app.destroy()
    assert csv_type == "csv" and xlsx_type == "xlsx"
    assert csv_count == expected and xlsx_count == expected
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        assert reader.fieldnames == db.EXPORT_FIELDS
        csv_rows = sum(1 for _ in reader)
    workbook = load_workbook(xlsx_path, read_only=True)
    sheet = workbook["Filtered Contacts"]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    xlsx_rows = sum(1 for _ in rows)
    workbook.close()
    assert headers == db.EXPORT_FIELDS
    assert csv_rows == expected and xlsx_rows == expected
    print({"status": "ok", "filtered_industry": "Medical NPI", "expected": expected, "csv_rows": csv_rows, "xlsx_rows": xlsx_rows, "columns": len(db.EXPORT_FIELDS)})
