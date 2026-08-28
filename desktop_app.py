from __future__ import annotations

import csv
import json
import os
import re
import tempfile
import threading
from datetime import datetime, timezone
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook

import clearout_client
import db
from importers import iter_rows, preview_file, supported_file
from dnc_client import DEFAULT_DNC_URL, DNCServiceError, check_service, clean_with_hosted_service

MAPPING_FIELDS = [
    ("name", "Name / Business name"), ("first_name", "First name"), ("last_name", "Last name"),
    ("email", "Email"), ("address_line_1", "Address line 1"), ("address_line_2", "Address line 2"),
    ("city", "City"), ("state", "State"), ("postal_code", "Postal / ZIP code"), ("country", "Country"),
    ("phone", "Phone"), ("fax", "Fax"), ("website", "Website"), ("original_industry", "Original industry"),
    ("alternate_categories", "Alternate categories"), ("credentials", "Credentials"),
    ("specialty", "Specialty"), ("rating", "Rating"), ("reviews", "Reviews"), ("npi", "NPI"),
    ("enumeration_date", "Enumeration date"), ("source_sheet", "Source sheet"),
]


def default_industry(filename: str) -> str:
    low = filename.lower()
    checks = [
        ("immigration.legal-canada", "Immigration Legal Canada"), ("mental_health", "Mental Health"),
        ("mental health", "Mental Health"), ("chiropr", "Chiropractor"), ("multi-spec", "Multi-Specialty"),
        ("family_medicine", "Family Medicine"), ("family medicine", "Family Medicine"),
        ("family practice", "Family Practice"), ("family_doctors", "Family Doctors"),
        ("internel", "Internal & Family Practice"), ("physicians", "Physicians"),
        ("phy-usa-npi", "Physicians"), ("med-npi", "Medical NPI"), ("medical-", "Medical"),
        ("accountant", "Accountant"), ("accounting", "Accountant"),
    ]
    for needle, label in checks:
        if needle in low:
            return label
    stem = Path(filename).stem
    stem = re.sub(r"\s*\([^)]*\)", "", stem)
    stem = re.sub(r"[-_]+", " ", stem)
    stem = re.sub(r"\d+", "", stem)
    return re.sub(r"\s+", " ", stem).strip().title() or "Imported"


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    normalized = {re.sub(r"[^a-z0-9]", "", h.lower()): h for h in headers}
    aliases = {
        "name": ["name", "businessname", "facilityname", "business", "companyname"],
        "first_name": ["firstname", "givenname"], "last_name": ["lastname", "surname", "familyname"],
        "email": ["email", "emailaddress", "mail"],
        "address_line_1": ["address", "streetaddress", "street", "firstlineaddress", "addressline1"],
        "address_line_2": ["secondlineaddress", "addressline2", "suite", "unit"],
        "city": ["city", "citytown", "cityname", "town"],
        "state": ["state", "stateprovince", "region", "sate", "province"],
        "postal_code": ["zip", "zipcode", "postalcode", "postcode"], "country": ["country", "countrycode"],
        "phone": ["phone", "telephone", "telephonenumber", "direct", "mobile"], "fax": ["fax"],
        "website": ["website", "url", "web"], "industry": ["industry", "occupation", "profession"],
        "original_industry": ["originalindustry"], "category": ["category", "primarycategory"],
        "alternate_categories": ["alternatecategories", "alternates", "subcategory"],
        "credentials": ["credentials", "credential"], "specialty": ["specialty", "speciality"],
        "rating": ["rating", "stars"], "reviews": ["reviews", "reviewcount", "numberofreviews"],
        "npi": ["npi", "npinumber"], "enumeration_date": ["enumerationdate", "enumeratedate"],
        "source": ["source", "provider", "datasource"], "source_type": ["sourcetype"],
        "source_sheet": ["sourcesheet", "sheet"], "source_category": ["sourcecategory"],
    }
    return {field: next((normalized[c] for c in candidates if c in normalized), "") for field, candidates in aliases.items()}


class ContactDirectoryDesktop(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Contact Directory")
        self.geometry("1500x900")
        self.minsize(1100, 700)
        db.init_db()
        self.pending_import = None
        self.pending_filter = None
        self.filter_output_path = None
        self._setup_style()
        self._build_ui()
        self.refresh_all()

    def _setup_style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Title.TLabel", font=("Segoe UI", 22, "bold"), foreground="#123349")
        style.configure("Subtitle.TLabel", font=("Segoe UI", 10), foreground="#687783")
        style.configure("Section.TLabel", font=("Segoe UI", 12, "bold"), foreground="#123349")
        style.configure("Accent.TButton", font=("Segoe UI", 10, "bold"))
        style.configure("Treeview", rowheight=28, font=("Segoe UI", 9))
        style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"))

    def _build_ui(self):
        header = ttk.Frame(self, padding=(20, 16))
        header.pack(fill="x")
        ttk.Label(header, text="Contact Directory", style="Title.TLabel").pack(anchor="w")
        ttk.Label(header, text="Native SQLite desktop application · imports, mapping, filtering, and audit metadata", style="Subtitle.TLabel").pack(anchor="w", pady=(3, 0))
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True, padx=14, pady=(0, 14))
        self.directory_tab = ttk.Frame(self.notebook, padding=12)
        self.import_tab = ttk.Frame(self.notebook, padding=12)
        self.filter_tab = ttk.Frame(self.notebook, padding=12)
        self.history_tab = ttk.Frame(self.notebook, padding=12)
        self.notebook.add(self.directory_tab, text="Directory")
        self.notebook.add(self.import_tab, text="Import & Map")
        self.notebook.add(self.filter_tab, text="Filter File")
        self.notebook.add(self.history_tab, text="Import History")
        self._build_directory()
        self._build_import()
        self._build_filter()
        self._build_history()

    def _build_directory(self):
        top = ttk.LabelFrame(self.directory_tab, text="Search and filters", padding=10)
        top.pack(fill="x", pady=(0, 10))
        self.filter_vars = {field: tk.StringVar() for field in ["q", "industry", "city", "state", "country", "category", "source", "source_type", "source_category", "source_file"]}
        labels = [("q", "Search"), ("industry", "Industry"), ("city", "City"), ("state", "State"), ("country", "Country"), ("category", "Category"), ("source", "Source"), ("source_type", "Source Type"), ("source_category", "Source category"), ("source_file", "Source file")]
        for index, (field, label) in enumerate(labels):
            row, col = divmod(index, 5)
            ttk.Label(top, text=label).grid(row=row * 2, column=col, sticky="w", padx=5, pady=(2, 0))
            if field == "q":
                widget = ttk.Entry(top, textvariable=self.filter_vars[field], width=27)
                widget.bind("<Return>", lambda _event: self.load_directory())
            else:
                widget = ttk.Combobox(top, textvariable=self.filter_vars[field], width=20, state="readonly")
                self.filter_widgets = getattr(self, "filter_widgets", {})
                self.filter_widgets[field] = widget
            widget.grid(row=row * 2 + 1, column=col, sticky="ew", padx=5, pady=(0, 5))
        for col in range(5):
            top.columnconfigure(col, weight=1)
        actions = ttk.Frame(top)
        action_row = ((len(labels) + 4) // 5) * 2
        actions.grid(row=action_row, column=0, columnspan=5, sticky="w", padx=5, pady=(5, 0))
        ttk.Button(actions, text="Apply filters", style="Accent.TButton", command=self.load_directory).pack(side="left", padx=(0, 6))
        ttk.Button(actions, text="Clear", command=self.clear_filters).pack(side="left", padx=(0, 6))
        ttk.Button(actions, text="Download CSV", command=self.export_csv).pack(side="left", padx=(0, 6))
        ttk.Button(actions, text="Download Excel", command=self.export_excel).pack(side="left")
        self.directory_status = ttk.Label(actions, text="")
        self.directory_status.pack(side="left", padx=15)

        tree_frame = ttk.Frame(self.directory_tab)
        tree_frame.pack(fill="both", expand=True)
        columns = [("name", "Name", 180), ("email", "Email", 190), ("address_line_1", "Address", 190), ("city", "City", 100), ("state", "State", 65), ("postal_code", "Postal", 80), ("phone", "Phone", 125), ("industry", "Industry", 125), ("category", "Category", 120), ("source", "Source", 130), ("uploaded_at", "Uploaded", 155), ("source_file", "File", 190)]
        self.directory_columns = [item[0] for item in columns]
        self.directory_tree = ttk.Treeview(tree_frame, columns=self.directory_columns, show="headings")
        for field, label, width in columns:
            self.directory_tree.heading(field, text=label)
            self.directory_tree.column(field, width=width, minwidth=60, anchor="w")
        yscroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.directory_tree.yview)
        xscroll = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.directory_tree.xview)
        self.directory_tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.directory_tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)

    def _build_import(self):
        scroll_host = ttk.Frame(self.import_tab)
        scroll_host.pack(fill="both", expand=True)
        self.import_scroll_canvas = tk.Canvas(scroll_host, highlightthickness=0, borderwidth=0)
        import_scrollbar = ttk.Scrollbar(scroll_host, orient="vertical", command=self.import_scroll_canvas.yview)
        self.import_scroll_canvas.configure(yscrollcommand=import_scrollbar.set)
        self.import_scroll_canvas.grid(row=0, column=0, sticky="nsew")
        import_scrollbar.grid(row=0, column=1, sticky="ns")
        scroll_host.rowconfigure(0, weight=1)
        scroll_host.columnconfigure(0, weight=1)
        import_content = ttk.Frame(self.import_scroll_canvas)
        import_window = self.import_scroll_canvas.create_window((0, 0), window=import_content, anchor="nw")
        import_content.bind("<Configure>", lambda _event: self.import_scroll_canvas.configure(scrollregion=self.import_scroll_canvas.bbox("all")))
        self.import_scroll_canvas.bind("<Configure>", lambda event: self.import_scroll_canvas.itemconfigure(import_window, width=event.width))
        self.import_scroll_canvas.bind("<Enter>", self._bind_import_scroll)
        self.import_scroll_canvas.bind("<Leave>", self._unbind_import_scroll)
        self.import_scroll_content = import_content

        self.import_file_var = tk.StringVar()
        self.import_date_var = tk.StringVar(value=self._now())
        self.import_header_vars = {}
        self.catalog_controls = {}
        top = ttk.LabelFrame(import_content, text="1. File and controlled import metadata", padding=12)
        top.pack(fill="x", pady=(0, 10))
        ttk.Label(top, text="File").grid(row=0, column=0, sticky="w", padx=5, pady=4)
        ttk.Entry(top, textvariable=self.import_file_var, width=68, state="readonly").grid(row=0, column=1, columnspan=4, sticky="ew", padx=5, pady=4)
        ttk.Button(top, text="Browse…", command=self.choose_import_file).grid(row=0, column=5, padx=5, pady=4)
        ttk.Label(top, text="These values come from reusable SQLite catalogs. Use Add New when the correct value does not exist.").grid(row=1, column=0, columnspan=6, sticky="w", padx=5, pady=(4, 10))
        self._catalog_selector(top, 2, 0, "source", "Source *", "sources", True)
        self._catalog_selector(top, 2, 2, "industry", "Industry *", "industries", True)
        self._catalog_selector(top, 2, 4, "source_type", "Source Type *", "source_types", True)
        self._catalog_selector(top, 4, 0, "source_category", "Source Category *", "source_categories", True)
        self._catalog_selector(top, 4, 2, "category", "Category", "categories", False)
        ttk.Label(top, text="Upload Date (automatic UTC)").grid(row=4, column=4, sticky="w", padx=5, pady=(2, 0))
        ttk.Entry(top, textvariable=self.import_date_var, width=26, state="readonly").grid(row=5, column=4, sticky="ew", padx=5, pady=(0, 4))
        ttk.Button(top, text="Read file and suggest mappings", command=self.read_import_file).grid(row=5, column=5, sticky="e", padx=5, pady=(0, 4))
        self.import_status = ttk.Label(top, text="Ready for a new import.", foreground="#687783")
        self.import_status.grid(row=6, column=0, columnspan=6, sticky="w", padx=5, pady=(4, 0))
        for col in range(6):
            top.columnconfigure(col, weight=1)

        self.import_mapping_frame = ttk.LabelFrame(import_content, text="2. Map source columns", padding=10)
        self.import_mapping_frame.pack(fill="x", pady=(0, 10))
        ttk.Label(self.import_mapping_frame, text="Controlled metadata above overrides same-named columns in the uploaded file. Original values remain available in raw_data.").pack(anchor="w")
        self.import_preview_frame = ttk.LabelFrame(import_content, text="3. Review sample and commit", padding=8)
        self.import_preview_frame.pack(fill="both", expand=True)
        ttk.Label(self.import_preview_frame, text="Choose a file above to load its sample preview.").pack(anchor="w")
        self.import_commit_button = ttk.Button(import_content, text="Validate and import mapped rows into SQLite", style="Accent.TButton", command=self.commit_import)
        self.import_commit_button.pack(anchor="e", pady=(10, 0))

    def _bind_import_scroll(self, _event=None):
        self.import_scroll_canvas.bind_all("<MouseWheel>", self._import_mousewheel)
        self.import_scroll_canvas.bind_all("<Button-4>", self._import_mousewheel)
        self.import_scroll_canvas.bind_all("<Button-5>", self._import_mousewheel)

    def _unbind_import_scroll(self, _event=None):
        self.import_scroll_canvas.unbind_all("<MouseWheel>")
        self.import_scroll_canvas.unbind_all("<Button-4>")
        self.import_scroll_canvas.unbind_all("<Button-5>")

    def _import_mousewheel(self, event):
        if getattr(event, "num", None) == 4:
            delta = -1
        elif getattr(event, "num", None) == 5:
            delta = 1
        else:
            delta = -1 * int(event.delta / 120) if event.delta else 0
        if delta:
            self.import_scroll_canvas.yview_scroll(delta, "units")

    def _catalog_selector(self, parent, row, col, key, label, table, required):
        frame = ttk.Frame(parent)
        frame.grid(row=row, column=col, columnspan=2, sticky="ew", padx=5, pady=3)
        frame.columnconfigure(0, weight=1)
        ttk.Label(frame, text=label).grid(row=0, column=0, columnspan=2, sticky="w")
        var = tk.StringVar()
        combo = ttk.Combobox(frame, textvariable=var, state="readonly", width=24)
        combo.grid(row=1, column=0, sticky="ew", padx=(0, 4))
        new_var = tk.StringVar()
        ttk.Entry(frame, textvariable=new_var, width=18).grid(row=1, column=1, sticky="ew", padx=(0, 4))
        ttk.Button(frame, text="Add New", command=lambda k=key: self._add_catalog_option(k)).grid(row=1, column=2, sticky="e")
        ttk.Label(frame, text="Type a new value, then select Add New." if required else "Optional; existing values are reusable.", foreground="#687783").grid(row=2, column=0, columnspan=3, sticky="w")
        self.catalog_controls[key] = {"table": table, "var": var, "combo": combo, "new_var": new_var, "required": required}

    def _add_catalog_option(self, key):
        control = self.catalog_controls[key]
        value = control["new_var"].get().strip()
        if not value:
            messagebox.showwarning("Add new value", "Enter a value first.")
            return
        try:
            db.add_option(control["table"], value)
            control["var"].set(value)
            control["new_var"].set("")
            self.refresh_import_options()
            self._set_import_status(f"Added and selected {value} in {key.replace('_', ' ').title()}.")
        except Exception as exc:
            messagebox.showerror("Add new value", str(exc))
    def _build_filter(self):
        self.filter_file_var = tk.StringVar()
        self.filter_phone_var = tk.StringVar()
        self.filter_email_var = tk.StringVar()
        self.filter_match_phone = tk.BooleanVar(value=True)
        self.filter_match_email = tk.BooleanVar(value=True)
        self.filter_run_hosted = tk.BooleanVar(value=False)
        self.filter_scrub_tcpa = tk.BooleanVar(value=True)
        self.dnc_url_var = tk.StringVar(value=DEFAULT_DNC_URL)
        self.clearout_token_var = tk.StringVar(value=os.environ.get("CLEAROUT_API_TOKEN", ""))
        self.clearout_country_var = tk.StringVar(value="us")
        top = ttk.LabelFrame(self.filter_tab, text="1. Select file", padding=10)
        top.pack(fill="x", pady=(0, 10))
        ttk.Entry(top, textvariable=self.filter_file_var, width=80, state="readonly").grid(row=0, column=0, sticky="ew", padx=5, pady=5)
        ttk.Button(top, text="Browse…", command=self.choose_filter_file).grid(row=0, column=1, padx=5, pady=5)
        ttk.Label(top, text="The database is never changed by this operation.").grid(row=1, column=0, columnspan=2, sticky="w", padx=5, pady=5)
        top.columnconfigure(0, weight=1)
        match = ttk.LabelFrame(self.filter_tab, text="2. Match rules and header mapping", padding=10)
        match.pack(fill="x", pady=(0, 10))
        ttk.Checkbutton(match, text="Match by Phone", variable=self.filter_match_phone).grid(row=0, column=0, sticky="w", padx=5, pady=5)
        ttk.Checkbutton(match, text="Match by Email", variable=self.filter_match_email).grid(row=0, column=2, sticky="w", padx=5, pady=5)
        ttk.Label(match, text="Phone column").grid(row=1, column=0, sticky="w", padx=5, pady=5)
        self.filter_phone_combo = ttk.Combobox(match, textvariable=self.filter_phone_var, width=30, state="readonly")
        self.filter_phone_combo.grid(row=1, column=1, sticky="w", padx=5, pady=5)
        ttk.Label(match, text="Email column").grid(row=1, column=2, sticky="w", padx=5, pady=5)
        self.filter_email_combo = ttk.Combobox(match, textvariable=self.filter_email_var, width=30, state="readonly")
        self.filter_email_combo.grid(row=1, column=3, sticky="w", padx=5, pady=5)
        ttk.Button(match, text="Compare and create clean file", style="Accent.TButton", command=self.run_filter).grid(row=2, column=0, columnspan=4, sticky="w", padx=5, pady=8)
        self.filter_status = ttk.Label(match, text="Select a file to load its headers.")
        self.filter_status.grid(row=3, column=0, columnspan=4, sticky="w", padx=5)

        hosted = ttk.LabelFrame(self.filter_tab, text="3. Optional hosted DNC / TCPA pass", padding=10)
        hosted.pack(fill="x", pady=(0, 10))
        ttk.Checkbutton(hosted, text="After local database filtering, run the remaining CSV through the hosted DNC cleaner", variable=self.filter_run_hosted).grid(row=0, column=0, columnspan=4, sticky="w", padx=5, pady=4)
        ttk.Checkbutton(hosted, text="Scrub phones against the hosted TCPA Litigator List", variable=self.filter_scrub_tcpa).grid(row=1, column=0, columnspan=2, sticky="w", padx=5, pady=4)
        ttk.Label(hosted, text="Cleaner URL").grid(row=1, column=2, sticky="e", padx=5, pady=4)
        ttk.Entry(hosted, textvariable=self.dnc_url_var, width=58).grid(row=1, column=3, sticky="ew", padx=5, pady=4)
        ttk.Button(hosted, text="Check service", command=self.check_dnc_service).grid(row=2, column=0, sticky="w", padx=5, pady=4)
        self.dnc_service_status = ttk.Label(hosted, text="Optional. If enabled, the remaining file is uploaded to this hosted service.", foreground="#687783")
        self.dnc_service_status.grid(row=2, column=1, columnspan=3, sticky="w", padx=5, pady=4)
        hosted.columnconfigure(3, weight=1)

        clearout = ttk.LabelFrame(self.filter_tab, text="4. Clearout phone validation (optional)", padding=10)
        clearout.pack(fill="x", pady=(0, 10))
        ttk.Label(clearout, text="API token").grid(row=0, column=0, sticky="w", padx=5, pady=4)
        ttk.Entry(clearout, textvariable=self.clearout_token_var, width=40, show="*").grid(row=0, column=1, sticky="ew", padx=5, pady=4)
        ttk.Label(clearout, text="Country code").grid(row=0, column=2, sticky="e", padx=5, pady=4)
        ttk.Entry(clearout, textvariable=self.clearout_country_var, width=8).grid(row=0, column=3, sticky="w", padx=5, pady=4)
        ttk.Label(clearout, text="After filtering, use the \"Send to Clearout\" button in the completion dialog to upload the output file directly for bulk phone validation.", foreground="#687783").grid(row=1, column=0, columnspan=4, sticky="w", padx=5, pady=(4, 0))
        clearout.columnconfigure(1, weight=1)

        self.filter_preview_frame = ttk.LabelFrame(self.filter_tab, text="5. Preview", padding=8)
        self.filter_preview_frame.pack(fill="both", expand=True)
        ttk.Label(self.filter_preview_frame, text="The preview will appear after selecting a file.").pack(anchor="w")

    def _build_history(self):
        toolbar = ttk.Frame(self.history_tab)
        toolbar.pack(fill="x", pady=(0, 8))
        ttk.Button(toolbar, text="Refresh", command=self.refresh_history).pack(anchor="e")
        frame = ttk.Frame(self.history_tab)
        frame.pack(fill="both", expand=True)
        cols = ["original_filename", "uploaded_at", "source", "industry", "source_type", "source_category", "category", "imported_at", "rows_imported", "rows_failed"]
        headings = {"original_filename": "File", "uploaded_at": "Uploaded date", "source": "Source", "industry": "Industry", "source_type": "Source Type", "source_category": "Source category", "category": "Category", "imported_at": "Imported at", "rows_imported": "Rows", "rows_failed": "Failed"}
        self.history_tree = ttk.Treeview(frame, columns=cols, show="headings")
        for col in cols:
            self.history_tree.heading(col, text=headings[col])
            self.history_tree.column(col, width=150 if col not in {"rows_imported", "rows_failed"} else 80)
        yscroll = ttk.Scrollbar(frame, orient="vertical", command=self.history_tree.yview)
        self.history_tree.configure(yscrollcommand=yscroll.set)
        self.history_tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)

    def _now(self):
        return datetime.now(timezone.utc).isoformat(timespec="seconds")

    def refresh_all(self):
        self.refresh_filter_options()
        self.refresh_directory()
        self.refresh_import_options()
        self.refresh_history()

    def refresh_filter_options(self):
        for field, widget in getattr(self, "filter_widgets", {}).items():
            widget["values"] = db.distinct_values(field)

    def refresh_import_options(self):
        if not hasattr(self, "catalog_controls"):
            return
        for key, control in self.catalog_controls.items():
            control["combo"]["values"] = db.option_values(control["table"])

    def _set_import_status(self, text):
        if hasattr(self, "import_status"):
            self.import_status.config(text=text)

    def _set_catalog_value(self, key, value):
        if key in self.catalog_controls and value:
            control = self.catalog_controls[key]
            if value not in control["combo"]["values"]:
                db.add_option(control["table"], value)
                self.refresh_import_options()
            control["var"].set(value)

    def clear_filters(self):
        for var in self.filter_vars.values():
            var.set("")
        self.load_directory()

    def load_directory(self):
        filters = {field: var.get() for field, var in self.filter_vars.items()}
        rows, total = db.query_contacts(filters, page=1, per_page=200)
        for item in self.directory_tree.get_children():
            self.directory_tree.delete(item)
        for row in rows:
            self.directory_tree.insert("", "end", values=[row[field] or "" for field in self.directory_columns])
        self.directory_status.config(text=f"Showing {len(rows):,} of {total:,} matching records")

    def _current_directory_filters(self):
        return {field: var.get() for field, var in self.filter_vars.items()}

    def export_csv(self):
        output = filedialog.asksaveasfilename(defaultextension=".csv", initialfile="filtered_contacts.csv", filetypes=[("CSV", "*.csv")])
        if not output:
            return
        self._start_export(Path(output), "csv")

    def export_excel(self):
        output = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="filtered_contacts.xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not output:
            return
        self._start_export(Path(output), "xlsx")

    def _start_export(self, output_path, file_type):
        filters = self._current_directory_filters()
        self.directory_status.config(text=f"Exporting filtered records to {output_path.name}…")
        self._run_background(lambda: self._export_rows(output_path, file_type, filters), self._export_finished)

    def _export_rows(self, output_path, file_type, filters):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        headers = list(db.EXPORT_FIELDS)
        count = 0
        if file_type == "csv":
            with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
                writer.writeheader()
                for row in db.iter_filtered_contacts(filters):
                    writer.writerow({field: row[field] or "" for field in headers})
                    count += 1
        else:
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("Filtered Contacts")
            sheet.append(headers)
            for row in db.iter_filtered_contacts(filters):
                sheet.append([row[field] or "" for field in headers])
                count += 1
            workbook.save(output_path)
        return str(output_path), count, file_type

    def _export_finished(self, result):
        output_path, count, file_type = result
        self.directory_status.config(text=f"Exported {count:,} filtered records to {Path(output_path).name}")
        messagebox.showinfo("Export complete", f"Exported {count:,} records to:\n{output_path}")

    def choose_import_file(self):
        path = filedialog.askopenfilename(filetypes=[("CSV or Excel", "*.csv *.xlsx *.xls"), ("All files", "*.*")])
        if path:
            self.import_file_var.set(path)
            self._set_catalog_value("industry", default_industry(Path(path).name))
            self.import_date_var.set(self._now())
            self._set_import_status(f"Selected {Path(path).name}. Review metadata and mappings before import.")
            self.read_import_file()

    def read_import_file(self):
        path = Path(self.import_file_var.get())
        if not path.exists() or not supported_file(path.name):
            messagebox.showerror("Import file", "Choose a valid CSV, XLSX, or XLS file first.")
            return
        try:
            headers, sample, sheets = preview_file(path)
        except Exception as exc:
            messagebox.showerror("Import file", str(exc))
            return
        self.pending_import = {"path": path, "headers": headers, "sample": sample, "sheets": sheets, "filename": path.name}
        self._render_import_mapping(headers)
        self._render_preview(self.import_preview_frame, headers, sample)

    def _clear_frame(self, frame):
        for child in frame.winfo_children():
            child.destroy()

    def _render_import_mapping(self, headers):
        self._clear_frame(self.import_mapping_frame)
        suggestions = suggest_mapping(headers)
        self.import_header_vars = {}
        for index, (field, label) in enumerate(MAPPING_FIELDS):
            row, col = divmod(index, 4)
            ttk.Label(self.import_mapping_frame, text=label).grid(row=row * 2, column=col, sticky="w", padx=5, pady=(2, 0))
            var = tk.StringVar(value=suggestions.get(field, ""))
            self.import_header_vars[field] = var
            combo = ttk.Combobox(self.import_mapping_frame, textvariable=var, values=[""] + headers, width=24)
            combo.grid(row=row * 2 + 1, column=col, sticky="ew", padx=5, pady=(0, 5))
        for col in range(4):
            self.import_mapping_frame.columnconfigure(col, weight=1)
        note = f"Loaded {len(headers)} headers" + (f" from sheets: {', '.join(self.pending_import['sheets'])}" if self.pending_import.get("sheets") else "")
        ttk.Label(self.import_mapping_frame, text=note).grid(row=((len(MAPPING_FIELDS) + 3) // 4) * 2, column=0, columnspan=4, sticky="w", padx=5, pady=(5, 0))

    def _render_preview(self, frame, headers, sample):
        self._clear_frame(frame)
        preview = ttk.Frame(frame)
        preview.pack(fill="both", expand=True)
        tree = ttk.Treeview(preview, columns=headers, show="headings", height=8)
        for header in headers:
            tree.heading(header, text=header)
            tree.column(header, width=150, minwidth=70)
        for row in sample:
            tree.insert("", "end", values=[row.get(header, "") for header in headers])
        yscroll = ttk.Scrollbar(preview, orient="vertical", command=tree.yview)
        xscroll = ttk.Scrollbar(preview, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        preview.rowconfigure(0, weight=1)
        preview.columnconfigure(0, weight=1)

    def commit_import(self):
        if not self.pending_import:
            messagebox.showerror("Import", "Choose and read a file first.")
            return
        values = {key: control["var"].get().strip() for key, control in self.catalog_controls.items()}
        missing = [key.replace("_", " ").title() for key, control in self.catalog_controls.items() if control["required"] and not values[key]]
        if missing:
            messagebox.showerror("Import validation", "Please select or add: " + ", ".join(missing))
            return
        mapping = {field: var.get() for field, var in self.import_header_vars.items()}
        if not any(mapping.get(field) for field in ("name", "first_name", "last_name", "email", "phone")):
            messagebox.showerror("Import validation", "Map at least one identifier: Name, First/Last Name, Email, or Phone.")
            return
        metadata = {"source": values["source"], "source_category": values["source_category"], "category": values.get("category", ""), "source_type": values["source_type"], "industry": values["industry"], "uploaded_at": self.import_date_var.get() or self._now(), "original_filename": self.pending_import["filename"]}
        self._set_import_status("Validated. Importing rows in the background…")
        self._run_background(lambda: db.insert_contacts(iter_rows(self.pending_import["path"]), mapping, metadata), self._import_finished)

    def _import_finished(self, result):
        imported, failed, errors = result
        self.pending_import = None
        message = f"Imported {imported:,} rows."
        if failed:
            message += f" Failed: {failed:,}."
        if errors:
            message += " " + " | ".join(errors[:2])
        messagebox.showinfo("Import complete", message)
        self._set_import_status(message)
        self.refresh_all()

    def choose_filter_file(self):
        path = filedialog.askopenfilename(filetypes=[("CSV or Excel", "*.csv *.xlsx *.xls"), ("All files", "*.*")])
        if not path:
            return
        path = Path(path)
        try:
            headers, sample, sheets = preview_file(path)
        except Exception as exc:
            messagebox.showerror("Filter file", str(exc))
            return
        self.pending_filter = {"path": path, "headers": headers, "sample": sample, "sheets": sheets, "filename": path.name}
        self.filter_phone_combo["values"] = headers
        self.filter_email_combo["values"] = headers
        suggestions = suggest_mapping(headers)
        self.filter_phone_var.set(suggestions.get("phone", ""))
        self.filter_email_var.set(suggestions.get("email", ""))
        self.filter_status.config(text=f"Loaded {len(headers)} headers. Select match rules and compare.")
        self._render_preview(self.filter_preview_frame, headers, sample)

    def run_filter(self):
        if not self.pending_filter:
            messagebox.showerror("Filter file", "Choose a file first.")
            return
        fields = []
        mapping = {}
        if self.filter_match_phone.get() and self.filter_phone_var.get():
            fields.append("phone")
            mapping["phone"] = self.filter_phone_var.get()
        if self.filter_match_email.get() and self.filter_email_var.get():
            fields.append("email")
            mapping["email"] = self.filter_email_var.get()
        if not fields:
            messagebox.showerror("Filter file", "Select at least one match rule and map its header.")
            return
        output = filedialog.asksaveasfilename(defaultextension=".csv", initialfile=f"filtered_{self.pending_filter['path'].stem}.csv", filetypes=[("CSV", "*.csv")])
        if not output:
            return
        run_hosted = self.filter_run_hosted.get()
        if run_hosted and "phone" not in fields:
            messagebox.showerror("Hosted DNC pass", "Map and enable Phone matching before using the hosted phone DNC/TCPA cleaner.")
            return
        service_url = self.dnc_url_var.get().strip() or DEFAULT_DNC_URL
        scrub_tcpa = self.filter_scrub_tcpa.get()
        self.filter_status.config(text="Filtering against the local database…" + (" Hosted DNC pass will follow." if run_hosted else ""))
        self._run_background(lambda: self._filter_rows(self.pending_filter["path"], output, mapping, fields, run_hosted, service_url, scrub_tcpa), self._filter_finished)

    def check_dnc_service(self):
        service_url = self.dnc_url_var.get().strip() or DEFAULT_DNC_URL
        self.dnc_service_status.config(text="Checking hosted cleaner…")
        self._run_background(lambda: check_service(service_url), self._dnc_check_finished)

    def _dnc_check_finished(self, result):
        ok, text = result
        self.dnc_service_status.config(text=text, foreground="#1d7a46" if ok else "#a23b3b")

    def _filter_rows(self, input_path, output_path, mapping, fields, run_hosted=False, service_url=DEFAULT_DNC_URL, scrub_tcpa=True):
        existing = db.existing_match_sets(fields)
        kept = removed = 0
        writer = None
        normalized_phone_rows = 0
        final_output = Path(output_path)
        temporary_dir = None
        local_output = final_output
        if run_hosted:
            temporary_dir = Path(tempfile.mkdtemp(prefix="contact-directory-local-filter-"))
            local_output = temporary_dir / "local_database_filtered.csv"
        with local_output.open("w", encoding="utf-8-sig", newline="") as output:
            for row in iter_rows(input_path):
                if "phone" in mapping:
                    phone_header = mapping["phone"]
                    raw_phone = row.get(phone_header, "")
                    canonical = db.canonical_phone(raw_phone)
                    if str(raw_phone or "").strip() != canonical:
                        normalized_phone_rows += 1
                    if phone_header:
                        row[phone_header] = canonical
                if writer is None:
                    headers = list(row.keys())
                    writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore", lineterminator="\n")
                    writer.writeheader()
                matches = set()
                if "email" in fields:
                    value = db.normalize_email(row.get(mapping["email"], ""))
                    if value:
                        matches.add(value in existing["email"])
                if "phone" in fields:
                    matches.add(bool(db.phone_keys(row.get(mapping["phone"], "")) & existing["phone"]))
                if any(matches):
                    removed += 1
                else:
                    writer.writerow(row)
                    kept += 1
        hosted_removed = 0
        hosted_standardized = 0
        final_kept = kept
        dnc_removed_path = None
        summary_path = None
        if run_hosted:
            try:
                hosted_result = clean_with_hosted_service(
                    local_output,
                    final_output,
                    phone_header=mapping.get("phone"),
                    service_url=service_url,
                    scrub_tcpa=scrub_tcpa,
                    progress=lambda text: self.after(0, lambda t=text: self.filter_status.config(text=t)),
                )
                final_kept = int(hosted_result["rows"])
                hosted_removed = max(0, kept - final_kept)
                hosted_standardized = int(hosted_result.get("standardized_phones", 0))
                dnc_removed_path = hosted_result.get("dnc_removed_path")
                summary_path = hosted_result.get("summary_path")
            finally:
                if temporary_dir:
                    import shutil
                    shutil.rmtree(temporary_dir, ignore_errors=True)
        return final_output, final_kept, removed + hosted_removed, normalized_phone_rows + hosted_standardized, run_hosted, dnc_removed_path, summary_path

    def _open_path(self, path):
        try:
            os.startfile(path)  # noqa: S606 - Windows-only desktop app
        except Exception as exc:
            messagebox.showerror("Open", f"Could not open:\n{path}\n\n{exc}")

    def _filter_finished(self, result):
        output_path, kept, removed, normalized_phone_rows, hosted_used, dnc_removed_path, summary_path = result
        hosted_note = " Hosted DNC/TCPA cleaning was applied after local database filtering." if hosted_used else ""
        self.filter_status.config(text=f"Complete: kept {kept:,}, removed {removed:,}, phones standardized {normalized_phone_rows:,}." + (" Hosted pass complete." if hosted_used else ""))

        dialog = tk.Toplevel(self)
        dialog.title("Filter complete")
        dialog.resizable(False, False)
        body = ttk.Frame(dialog, padding=12)
        body.grid(row=0, column=0, sticky="nsew")
        message = (
            f"Clean file saved to:\n{output_path}\n\n"
            f"Rows kept: {kept:,}\nRows removed: {removed:,}\n\n"
            f"Phone values standardized: {normalized_phone_rows:,}\nThe database was not changed.{hosted_note}"
        )
        ttk.Label(body, text=message, justify="left").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 10))

        row = 1
        ttk.Button(body, text="Open output file", command=lambda: self._open_path(output_path)).grid(row=row, column=0, sticky="w", padx=(0, 6))
        ttk.Button(body, text="Open containing folder", command=lambda: self._open_path(Path(output_path).parent)).grid(row=row, column=1, sticky="w", padx=(0, 6))
        clearout_button = ttk.Button(body, text="Send to Clearout", style="Accent.TButton")
        clearout_button.grid(row=row, column=2, sticky="w", padx=(0, 6))
        row += 1
        clearout_status = ttk.Label(body, text="", foreground="#687783", wraplength=440, justify="left")
        clearout_status.grid(row=row, column=0, columnspan=3, sticky="w", pady=(4, 0))
        clearout_button.configure(command=lambda: self._send_to_clearout(output_path, clearout_button, clearout_status))
        row += 1
        if dnc_removed_path:
            ttk.Label(body, text="Rows removed by the hosted DNC/TCPA pass were saved separately:").grid(row=row, column=0, columnspan=3, sticky="w", pady=(10, 2))
            row += 1
            ttk.Button(body, text="Open DNC-removed rows", command=lambda: self._open_path(dnc_removed_path)).grid(row=row, column=0, sticky="w", padx=(0, 6))
            row += 1
        if summary_path:
            ttk.Button(body, text="Open cleaning summary", command=lambda: self._open_path(summary_path)).grid(row=row, column=0, sticky="w", padx=(0, 6))
            row += 1
        ttk.Button(body, text="Close", command=dialog.destroy).grid(row=row, column=0, sticky="w", pady=(10, 0))
        dialog.transient(self)
        dialog.grab_set()

    def refresh_directory(self):
        if hasattr(self, "directory_tree"):
            self.load_directory()

    def refresh_history(self):
        if not hasattr(self, "history_tree"):
            return
        for item in self.history_tree.get_children():
            self.history_tree.delete(item)
        for row in db.recent_imports(100):
            self.history_tree.insert("", "end", values=[row[col] for col in ["original_filename", "uploaded_at", "source", "industry", "source_type", "source_category", "category", "imported_at", "rows_imported", "rows_failed"]])

    def _run_background(self, function, callback, on_error=None):
        def worker():
            try:
                result = function()
                self.after(0, lambda: callback(result))
            except Exception as exc:
                if on_error:
                    self.after(0, lambda: on_error(exc))
                else:
                    self.after(0, lambda: messagebox.showerror("Operation failed", str(exc)))
        threading.Thread(target=worker, daemon=True).start()

    def _send_to_clearout(self, output_path, button, status_label):
        token = self.clearout_token_var.get().strip()
        if not token:
            messagebox.showerror("Send to Clearout", "Enter a Clearout API token in the Filter File tab first.")
            return
        country_code = self.clearout_country_var.get().strip() or "us"
        button.state(["disabled"])
        status_label.config(text="Sending output file to Clearout…", foreground="#687783")

        def finished(result):
            button.state(["!disabled"])
            list_id = (result.get("data") or {}).get("list_id", "")
            status_label.config(
                text=f"Clearout accepted the file (list ID: {list_id})." if list_id else "Clearout responded.",
                foreground="#1d7a46",
            )
            messagebox.showinfo("Clearout response", json.dumps(result, indent=2))

        def failed(exc):
            button.state(["!disabled"])
            status_label.config(text="Clearout request failed.", foreground="#a23b3b")
            messagebox.showerror("Send to Clearout", str(exc))

        self._run_background(
            lambda: clearout_client.send_bulk_validation(output_path, token, country_code=country_code),
            finished,
            failed,
        )


if __name__ == "__main__":
    ContactDirectoryDesktop().mainloop()
